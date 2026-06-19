import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

def create_project_documentation():
    # Initialize document
    doc = Document()
    
    # Title Page
    doc.add_heading('Project Documentation', 0)
    title = doc.add_heading('AI Interview Platform', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('\n\n\n')
    
    # Introduction
    doc.add_heading('1. Introduction', level=1)
    doc.add_paragraph(
        "The AI Interview Platform is a comprehensive web-based application designed to streamline the recruitment process. "
        "It leverages Artificial Intelligence to conduct, analyze, and manage interviews autonomously. The platform bridges "
        "the gap between HR professionals and candidates by automating initial screening and question generation, allowing "
        "companies to evaluate candidates efficiently and accurately at scale."
    )
    
    # Objectives
    doc.add_heading('2. Objectives', level=1)
    doc.add_paragraph(
        "• Automate the interview scheduling and screening process.\n"
        "• Generate dynamic, AI-driven technical and behavioral questions.\n"
        "• Provide a seamless and interactive interface for candidates to take live interviews.\n"
        "• Assist HR and Admins in managing job postings, candidate evaluations, and approval queues."
    )
    
    # Modules
    doc.add_heading('3. Key Modules', level=1)
    
    doc.add_heading('3.1 Admin Panel', level=2)
    doc.add_paragraph(
        "The Admin module serves as the central control for the entire platform. Admins have the authority to manage the system, "
        "oversee HR activities, and handle job approvals."
    )
    p = doc.add_paragraph()
    p.add_run("Features:").bold = True
    doc.add_paragraph("• Dashboard for platform overview and analytics.", style='List Bullet')
    doc.add_paragraph("• Job Approval Queue for vetting HR job postings.", style='List Bullet')
    doc.add_paragraph("• User and role management.", style='List Bullet')
    doc.add_paragraph("• Security and system configurations.", style='List Bullet')
    
    doc.add_heading('3.2 HR Module', level=2)
    doc.add_paragraph(
        "The HR module enables recruiters to manage hiring campaigns effectively."
    )
    p = doc.add_paragraph()
    p.add_run("Features:").bold = True
    doc.add_paragraph("• Create and manage job postings.", style='List Bullet')
    doc.add_paragraph("• Schedule and track interviews.", style='List Bullet')
    doc.add_paragraph("• Review candidate performance and AI-generated scores.", style='List Bullet')
    doc.add_paragraph("• Automate email communications and notifications.", style='List Bullet')
    
    doc.add_heading('3.3 Candidate Module', level=2)
    doc.add_paragraph(
        "Provides candidates with a dedicated portal to view opportunities and participate in AI-driven interviews."
    )
    p = doc.add_paragraph()
    p.add_run("Features:").bold = True
    doc.add_paragraph("• View assigned interviews and job details.", style='List Bullet')
    doc.add_paragraph("• Participate in live, AI-conducted interview sessions.", style='List Bullet')
    doc.add_paragraph("• Real-time video/audio and text-based question answering.", style='List Bullet')
    doc.add_paragraph("• Instant feedback and completion notifications.", style='List Bullet')

    # Core Technologies
    doc.add_heading('4. Core Technologies (Tech Stack)', level=1)
    doc.add_paragraph("• Backend Framework: Django (Python)", style='List Bullet')
    doc.add_paragraph("• Frontend: HTML5, CSS3, JavaScript", style='List Bullet')
    doc.add_paragraph("• Database: SQLite (configurable for PostgreSQL/MySQL)", style='List Bullet')
    doc.add_paragraph("• AI Integration: Custom question generators and natural language processing (NLP).", style='List Bullet')
    doc.add_paragraph("• Authentication: Secure Role-Based Access Control (RBAC).", style='List Bullet')

    # Architecture Overview
    doc.add_heading('5. Architecture Overview', level=1)
    doc.add_paragraph(
        "The platform follows a robust Model-View-Template (MVT) architecture inherent to Django. "
        "The backend APIs securely communicate with the frontend to deliver dynamic content. "
        "AI utilities (like question generation) are modularized within the core logic, ensuring that they can be "
        "scaled or swapped out without impacting the entire application."
    )

    # Future Scope
    doc.add_heading('6. Future Scope', level=1)
    doc.add_paragraph("• Integration with advanced voice-to-text models for real-time conversational interviews.", style='List Bullet')
    doc.add_paragraph("• Advanced analytics and predictive hiring insights for HR.", style='List Bullet')
    doc.add_paragraph("• Resume parsing and automatic candidate-to-job matching.", style='List Bullet')

    # Conclusion
    doc.add_heading('7. Conclusion', level=1)
    doc.add_paragraph(
        "The AI Interview Platform effectively modernizes recruitment. By abstracting the repetitive aspects of interviewing, "
        "it allows companies to focus on evaluating top talent faster and more fairly."
    )

    # Save Document
    desktop_path = os.path.join(os.path.expanduser("~"), "Desktop", "AI_Interview_Platform_Documentation.docx")
    doc.save(desktop_path)
    print(f"Document generated successfully at {desktop_path}")

if __name__ == "__main__":
    create_project_documentation()
